"""Build the Russian SignBridge project presentation as a PowerPoint deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("SignBridge_presentation_ru.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


BG = rgb("080B14")
SURFACE = rgb("111827")
SURFACE_2 = rgb("172033")
TEXT = rgb("F8FAFC")
MUTED = rgb("94A3B8")
CYAN = rgb("00D4FF")
VIOLET = rgb("8B5CF6")
GREEN = rgb("10B981")
ORANGE = rgb("F59E0B")
PINK = rgb("EC4899")
LINE = rgb("293548")
FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"


def add_shape(slide, shape_type, x, y, w, h, fill, line_fill=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_fill or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, text, x, y, w, h, size=16, color=TEXT, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_paragraphs(slide, items, x, y, w, h, size=15, color=TEXT, bullet_color=CYAN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.text = f"•  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
    return box


def base(slide, number, section="SIGNBRIDGE"):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 0.075, 7.5, CYAN)
    add_text(slide, section, 0.45, 7.08, 2.5, 0.18, 8, MUTED, bold=True)
    add_text(slide, f"{number:02d}", 12.35, 7.03, 0.45, 0.22, 9, MUTED, bold=True, align=PP_ALIGN.RIGHT)


def title(slide, kicker, heading, subtitle, number, section="ПРОЕКТ"):
    base(slide, number, section)
    add_text(slide, kicker.upper(), 0.55, 0.45, 5.0, 0.25, 10, CYAN, bold=True)
    add_text(slide, heading, 0.55, 0.85, 11.8, 0.82, 27, TEXT, bold=True, font=FONT_DISPLAY)
    add_text(slide, subtitle, 0.57, 1.78, 11.6, 0.45, 13, MUTED)


def card(slide, x, y, w, h, accent=None):
    shape = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, SURFACE, LINE, True)
    if accent:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, 0.06, h, accent)
    return shape


def pill(slide, text, x, y, w, color=CYAN):
    shape = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.31, SURFACE_2, color, True)
    add_text(slide, text, x + 0.08, y + 0.065, w - 0.16, 0.14, 9, color, bold=True, align=PP_ALIGN.CENTER)
    return shape


def arrow(slide, x1, y1, x2, y2, color=CYAN, width=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    return connector


def metric(slide, value, label, x, y, w, color):
    add_text(slide, value, x, y, w, 0.4, 24, color, bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y + 0.45, w, 0.35, 10, MUTED, align=PP_ALIGN.CENTER)


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 1, "ПРЕЗЕНТАЦИЯ")
add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.62, 0.62, 1.1, 1.1, SURFACE_2, CYAN, True)
add_text(slide, "SB", 0.72, 0.88, 0.9, 0.32, 23, CYAN, bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER)
add_text(slide, "SignBridge", 0.62, 2.05, 10.8, 0.8, 38, TEXT, bold=True, font=FONT_DISPLAY)
add_text(slide, "Локальная платформа для распознавания жестов и\nвизуализации коммуникации в реальном времени", 0.65, 2.94, 7.5, 0.72, 18, MUTED)
pill(slide, "MediaPipe + FastAPI + WebSocket", 0.65, 4.0, 2.72, CYAN)
pill(slide, "PWA и Canvas", 3.52, 4.0, 1.72, VIOLET)
pill(slide, "Local-first", 5.38, 4.0, 1.36, GREEN)
add_text(slide, "Техническая презентация проекта", 0.67, 6.35, 4.2, 0.25, 11, MUTED)

# Decorative hand nodes
for x, y, r, color in [(9.0, 1.42, 0.23, CYAN), (10.25, 1.1, 0.19, VIOLET), (11.4, 1.75, 0.24, GREEN),
                       (10.0, 2.68, 0.2, ORANGE), (11.95, 3.05, 0.16, PINK), (9.15, 3.55, 0.2, CYAN),
                       (10.85, 4.05, 0.24, VIOLET), (11.8, 5.02, 0.18, GREEN)]:
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, r, r, color)
for points in [(9.2, 1.63, 10.34, 1.28), (10.42, 1.29, 11.52, 1.92), (10.16, 2.85, 11.99, 3.18),
               (9.35, 3.67, 10.98, 4.18), (10.96, 4.22, 11.89, 5.14), (9.16, 1.65, 10.05, 2.78)]:
    arrow(slide, *points, color=LINE, width=2)
add_text(slide, "Жест → смысл → понятная коммуникация", 8.55, 5.78, 3.8, 0.3, 13, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 2. Context
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "01. Контекст", "Коммуникация должна быть доступной без посредников", "SignBridge исследует практический путь от видеопотока руки к понятному сообщению.", 2)
for i, (head, body, color) in enumerate([
    ("Барьер", "Жестовый язык требует визуального канала и не всегда доступен собеседнику.", PINK),
    ("Задержка", "Даже точное распознавание теряет ценность, когда видеопоток и контур руки отстают.", ORANGE),
    ("Доверие", "Пользователю нужны понятные статусы, предсказуемость и возможность видеть ход распознавания.", GREEN),
]):
    x = 0.65 + i * 4.14
    card(slide, x, 2.65, 3.72, 2.6, color)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.35, 3.02, 0.55, 0.55, color)
    add_text(slide, f"0{i + 1}", x + 0.35, 3.17, 0.55, 0.15, 10, BG, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, head, x + 0.35, 3.86, 2.8, 0.3, 18, TEXT, bold=True)
    add_text(slide, body, x + 0.35, 4.34, 2.9, 0.6, 12, MUTED)
add_text(slide, "Цель проекта: сохранить живое ощущение камеры и превратить движения руки в удобный двусторонний интерфейс.", 0.67, 5.9, 11.6, 0.38, 15, TEXT, bold=True)

# 3. Value flow
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "02. Решение", "Один поток, четыре понятных шага", "Пользователь видит происходящее, система распознаёт жест, а результат сразу превращается в коммуникационное действие.", 3)
steps = [
    ("01", "Камера", "Живой видеопоток и локальный контур 21 точки руки.", CYAN),
    ("02", "Распознавание", "Статические и динамические жесты в серверном движке.", VIOLET),
    ("03", "Смысл", "Перевод, накопление слова, история результатов.", GREEN),
    ("04", "Ответ", "Canvas-аватар и TTS-план для обратной коммуникации.", ORANGE),
]
for i, (num, head, body, color) in enumerate(steps):
    x = 0.62 + i * 3.13
    card(slide, x, 2.55, 2.72, 2.7, color)
    add_text(slide, num, x + 0.28, 2.87, 0.55, 0.24, 13, color, bold=True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.3, 3.42, 0.54, 0.54, color)
    add_text(slide, head, x + 0.28, 4.25, 2.1, 0.28, 16, TEXT, bold=True)
    add_text(slide, body, x + 0.28, 4.65, 2.1, 0.4, 11, MUTED)
    if i < 3:
        arrow(slide, x + 2.78, 3.88, x + 3.08, 3.88, color=CYAN)
add_text(slide, "Принцип UX: обработка не должна скрывать или подменять живое изображение с камеры.", 0.68, 5.95, 11.2, 0.35, 15, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 4. Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "03. Архитектура", "Лёгкий фронтенд, прозрачный API, локальный движок жестов", "Система разделяет визуальную обратную связь, распознавание и анимацию аватара.", 4)
layers = [
    ("Браузер / PWA", "HTML + CSS + Vanilla JS\nCanvas overlay · Service Worker · Web Speech", CYAN, 0.7),
    ("Два WebSocket-канала", "/ws/gesture: JPEG-кадры и результат\n/ws/avatar: текст, жестовый план, аудио", VIOLET, 3.95),
    ("FastAPI", "Валидация · ограничение размеров · CORS · health/readiness\nREST /api/v1 + статическая раздача", GREEN, 7.2),
    ("Gesture Engine", "MediaPipe Hand Landmarker · optional sklearn\nOpenCV fallback · dynamic motion buffer", ORANGE, 10.45),
]
for head, body, color, x in layers:
    card(slide, x, 2.62, 2.2, 2.4, color)
    add_text(slide, head, x + 0.22, 2.95, 1.75, 0.45, 13, TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.2, 3.7, 1.8, 0.75, 10, MUTED, align=PP_ALIGN.CENTER)
for x in [2.98, 6.23, 9.48]:
    arrow(slide, x, 3.82, x + 0.8, 3.82, color=CYAN, width=2)
add_text(slide, "Модели и обработка остаются на локальной машине; для базового сценария не нужен облачный inference API.", 0.72, 5.75, 11.5, 0.45, 14, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 5. Camera loop
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "04. Камера", "Два контура обработки вместо одного тяжёлого кадра", "Отрисовка нужна пользователю прямо сейчас; классификация может идти отдельным контролируемым потоком.", 5)
card(slide, 0.72, 2.55, 5.65, 3.4, CYAN)
add_text(slide, "Контур визуальной обратной связи", 1.05, 2.88, 4.95, 0.28, 17, TEXT, bold=True)
add_text(slide, "Browser MediaPipe Hand Landmarker", 1.05, 3.33, 4.8, 0.26, 12, CYAN, bold=True)
add_paragraphs(slide, ["Берёт текущий декодированный кадр video", "Рисует 21 точку и связи на canvas-overlay", "Ограничен до 15 FPS, чтобы не вытеснять распознавание"], 1.05, 3.85, 4.75, 1.25, 13)
metric(slide, "15 FPS", "локальный контур", 1.1, 5.2, 1.45, CYAN)
metric(slide, "21", "landmark-точка", 2.85, 5.2, 1.35, VIOLET)
metric(slide, "0", "JPEG-overlay", 4.5, 5.2, 1.25, GREEN)
card(slide, 6.85, 2.55, 5.65, 3.4, VIOLET)
add_text(slide, "Контур распознавания", 7.18, 2.88, 4.95, 0.28, 17, TEXT, bold=True)
add_text(slide, "WebSocket → FastAPI → GestureEngine", 7.18, 3.33, 4.8, 0.26, 12, VIOLET, bold=True)
add_paragraphs(slide, ["Сжимает и отправляет кадр только при свободном запросе", "Получает жест, уверенность, перевод и накопленное слово", "Не передаёт обратно аннотированное видео"], 7.18, 3.85, 4.75, 1.25, 13, bullet_color=VIOLET)
metric(slide, "640×480", "размер кадра", 7.25, 5.2, 1.5, VIOLET)
metric(slide, "80 ms", "интервал отправки", 9.02, 5.2, 1.5, ORANGE)
metric(slide, "1", "кадр in-flight", 10.85, 5.2, 1.25, GREEN)

# 6. Recognition pipeline
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "05. Распознавание", "Пайплайн объединяет landmarks, классификацию и движение", "Движок выбирает доступный режим и не теряет работоспособность при отсутствии модели.", 6)
pipeline = [
    ("JPEG", "decode + size limits", CYAN),
    ("RGB", "MediaPipe image", VIOLET),
    ("21×3", "x, y, z landmarks", GREEN),
    ("Motion", "buffer до 30 кадров", ORANGE),
    ("Result", "gesture + confidence", PINK),
]
for i, (head, body, color) in enumerate(pipeline):
    x = 0.62 + i * 2.5
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.55, 2.8, 1.22, 1.22, SURFACE_2, color)
    add_text(slide, head, x + 0.7, 3.15, 0.92, 0.2, 13, color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.16, 4.35, 2.0, 0.36, 10, MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(slide, x + 1.85, 3.42, x + 2.62, 3.42, color=LINE, width=2)
card(slide, 0.72, 5.2, 3.7, 0.82, VIOLET)
add_text(slide, "Динамические жесты", 1.0, 5.42, 1.7, 0.18, 12, TEXT, bold=True)
add_text(slide, "Правила по траектории wrist и последовательности landmarks", 2.45, 5.38, 1.65, 0.25, 9, MUTED, align=PP_ALIGN.RIGHT)
card(slide, 4.82, 5.2, 3.7, 0.82, GREEN)
add_text(slide, "Статические знаки", 5.12, 5.42, 1.5, 0.18, 12, TEXT, bold=True)
add_text(slide, "Optional sklearn classifier; fallback к эвристикам", 6.55, 5.38, 1.6, 0.25, 9, MUTED, align=PP_ALIGN.RIGHT)
card(slide, 8.92, 5.2, 3.58, 0.82, ORANGE)
add_text(slide, "Fallback", 9.22, 5.42, 0.8, 0.18, 12, TEXT, bold=True)
add_text(slide, "OpenCV: skin mask, contour, convexity defects", 10.0, 5.38, 2.1, 0.25, 9, MUTED, align=PP_ALIGN.RIGHT)

# 7. Latency decision
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "06. Оптимизация", "Почему контур раньше отставал и как это исправлено", "Ключевое решение: отделить отрисовку руки от возврата обработанного видеокадра.", 7)
card(slide, 0.72, 2.53, 5.7, 3.5, PINK)
add_text(slide, "Было: серверный JPEG-overlay", 1.05, 2.88, 4.95, 0.3, 17, TEXT, bold=True)
add_text(slide, "video → JPEG → WebSocket → inference → JPEG → canvas", 1.05, 3.38, 4.8, 0.24, 12, PINK, bold=True)
add_paragraphs(slide, ["Локальное видео скрывалось отстающим серверным снимком", "Каждый результат требовал дополнительного JPEG-кодирования", "Задержка визуально выглядела как медленная камера"], 1.05, 3.9, 4.75, 1.2, 13, bullet_color=PINK)
add_text(slide, "Проблема: связь между качеством inference и плавностью UI", 1.05, 5.43, 4.65, 0.3, 11, MUTED, bold=True)
card(slide, 6.9, 2.53, 5.6, 3.5, GREEN)
add_text(slide, "Теперь: локальный overlay + серверный результат", 7.23, 2.88, 4.85, 0.3, 17, TEXT, bold=True)
add_text(slide, "video → browser landmarks → canvas     |     frame → server → text", 7.23, 3.38, 4.8, 0.24, 12, GREEN, bold=True)
add_paragraphs(slide, ["Canvas отражает текущий видеокадр, а не прошлый ответ сервера", "В сеть уходит только кадр для распознавания, без annotated_frame", "Backpressure: один запрос в обработке, timeout на восстановление"], 7.23, 3.9, 4.75, 1.2, 13, bullet_color=GREEN)
add_text(slide, "Результат: визуальная отзывчивость и распознавание можно оптимизировать независимо", 7.23, 5.43, 4.65, 0.3, 11, MUTED, bold=True)

# 8. Frontend
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "07. Продуктовый интерфейс", "Не просто demo-экран, а рабочее пространство для диалога", "Фронтенд спроектирован вокруг повторяющегося действия: показать жест, проверить результат, передать смысл.", 8)
card(slide, 0.7, 2.45, 6.0, 3.65, CYAN)
add_text(slide, "Режим распознавания", 1.03, 2.78, 2.6, 0.25, 16, TEXT, bold=True)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.04, 3.25, 3.4, 2.2, rgb("050810"), LINE, True)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 2.55, 3.82, 0.45, 0.45, CYAN)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 3.22, 4.0, 0.32, 0.32, VIOLET)
arrow(slide, 2.85, 4.04, 3.37, 4.15, color=CYAN, width=2)
add_text(slide, "LIVE", 1.28, 3.55, 0.5, 0.16, 9, GREEN, bold=True)
add_text(slide, "Контур руки", 1.36, 5.1, 1.4, 0.18, 10, MUTED)
add_text(slide, "Распознанный жест", 4.82, 3.28, 1.45, 0.18, 11, MUTED, align=PP_ALIGN.CENTER)
add_text(slide, "HELLO", 4.76, 3.72, 1.58, 0.35, 20, CYAN, bold=True, align=PP_ALIGN.CENTER)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 4.98, 4.35, 1.15, 0.09, GREEN)
add_text(slide, "96% confidence", 4.78, 4.62, 1.55, 0.18, 10, MUTED, align=PP_ALIGN.CENTER)
add_text(slide, "Накопленное слово · история · языковые переключатели · TTS", 1.04, 5.72, 5.1, 0.18, 11, MUTED)
card(slide, 7.1, 2.45, 5.37, 3.65, VIOLET)
add_text(slide, "Режим аватара", 7.43, 2.78, 2.6, 0.25, 16, TEXT, bold=True)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.16, 3.28, 1.15, 1.15, SURFACE_2, VIOLET)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.ARC, 8.72, 4.25, 2.02, 1.15, SURFACE_2, VIOLET)
add_text(slide, "Text → gesture plan → canvas animation → optional audio", 7.55, 5.3, 4.37, 0.24, 11, MUTED, align=PP_ALIGN.CENTER)
add_text(slide, "Поддерживаемые языки интерфейса: RU · EN · UZ", 7.55, 5.7, 4.37, 0.18, 11, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 9. API surface
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "08. Интеграция", "API разделяет здоровье системы, перевод, голос и realtime", "Маршруты имеют версионирование /api/v1; WebSocket остаётся для интерактивных сценариев.", 9)
api_rows = [
    ("GET", "/api/v1/health", "Статус сервиса и режим распознавания", GREEN),
    ("GET", "/api/v1/vocabulary", "Доступные глоссы и словарь", CYAN),
    ("POST", "/api/v1/translate", "Проверяемый план жестовой анимации", VIOLET),
    ("POST", "/api/v1/tts", "План аватара и аудио", ORANGE),
    ("WS", "/ws/gesture", "Кадр → gesture, confidence, translation, word", PINK),
    ("WS", "/ws/avatar", "Текст → sequence + optional audio", GREEN),
]
for i, (method, path, description, color) in enumerate(api_rows):
    y = 2.35 + i * 0.63
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.78, y, 0.78, 0.34, SURFACE_2, color, True)
    add_text(slide, method, 0.84, y + 0.09, 0.65, 0.12, 9, color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, path, 1.9, y + 0.06, 3.25, 0.2, 12, TEXT, bold=True, font="Consolas")
    add_text(slide, description, 5.5, y + 0.06, 5.85, 0.2, 12, MUTED)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 11.83, y + 0.11, 0.12, 0.12, color)
add_text(slide, "Инженерный акцент: валидация payload, пределы размера кадров и WebSocket-origin policy снижают риск непредсказуемой нагрузки.", 0.8, 6.45, 11.6, 0.3, 13, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 10. Trust and deployment
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "09. Надёжность", "Local-first развёртывание и контролируемые границы", "Проект не позиционируется как сертифицированный переводчик: это технический прототип с прозрачными ограничениями.", 10)
trust = [
    ("Локальная модель", "Hand Landmarker загружается на машину; распознавание не зависит от внешнего inference API.", GREEN),
    ("Ограничения входа", "Размеры WebSocket-сообщений, JPEG-кадров и разрешения проверяются до обработки.", CYAN),
    ("Fallback", "При недоступности MediaPipe движок переключается на OpenCV-контуры и эвристики.", ORANGE),
    ("PWA", "Service Worker кэширует shell приложения для повторных запусков и offline-сценариев.", VIOLET),
]
for i, (head, body, color) in enumerate(trust):
    x = 0.7 + (i % 2) * 6.2
    y = 2.45 + (i // 2) * 1.75
    card(slide, x, y, 5.7, 1.28, color)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.35, y + 0.34, 0.5, 0.5, color)
    add_text(slide, head, x + 1.1, y + 0.29, 3.7, 0.2, 15, TEXT, bold=True)
    add_text(slide, body, x + 1.1, y + 0.63, 4.05, 0.3, 11, MUTED)
add_text(slide, "Deployment: Windows launcher → virtual environment → model download → Uvicorn → localhost", 0.82, 6.25, 11.6, 0.27, 13, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 11. Demo
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "10. Демонстрация", "Сценарий показа на 90 секунд", "Демонстрация построена так, чтобы аудитория видела и технологию, и пользовательскую ценность.", 11)
demo = [
    ("00–15 c", "Открыть localhost", "Показать статус готовности, PWA-интерфейс и выбор языка.", CYAN),
    ("15–35 c", "Включить камеру", "Контур руки появляется поверх живого видео без подмены кадра.", VIOLET),
    ("35–55 c", "Показать жест", "Уверенность, перевод и история обновляются в интерфейсе.", GREEN),
    ("55–75 c", "Собрать слово", "Накопление букв демонстрирует переход от жеста к сообщению.", ORANGE),
    ("75–90 c", "Ответ аватара", "Текст превращается в жестовый план и визуальную обратную коммуникацию.", PINK),
]
for i, (time, head, body, color) in enumerate(demo):
    y = 2.28 + i * 0.77
    add_text(slide, time, 0.78, y + 0.12, 1.0, 0.2, 11, color, bold=True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 2.02, y + 0.1, 0.3, 0.3, color)
    if i < 4:
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.17), Inches(y + 0.4), Inches(2.17), Inches(y + 0.84))
        connector.line.color.rgb = LINE
        connector.line.width = Pt(2)
    add_text(slide, head, 2.65, y + 0.04, 2.3, 0.22, 15, TEXT, bold=True)
    add_text(slide, body, 5.0, y + 0.04, 6.7, 0.3, 12, MUTED)

# 12. Quality and risks
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "11. Качество", "Что уже проверяется и что нельзя обещать", "Техническая честность важнее красивой демонстрации: у распознавания есть измеримые риски и границы применимости.", 12)
card(slide, 0.72, 2.48, 5.65, 3.25, GREEN)
add_text(slide, "Уже заложено", 1.05, 2.83, 3.8, 0.28, 18, TEXT, bold=True)
add_paragraphs(slide, ["Python compileall и JavaScript syntax checks", "Smoke-тесты API и планировщика жестовых последовательностей", "Health / readiness endpoints и явный recognition_mode", "Timeout/backpressure для кадра в обработке"], 1.05, 3.43, 4.75, 1.65, 13, bullet_color=GREEN)
card(slide, 6.85, 2.48, 5.65, 3.25, ORANGE)
add_text(slide, "Нужно валидировать перед пилотом", 7.18, 2.83, 4.4, 0.28, 18, TEXT, bold=True)
add_paragraphs(slide, ["Точность на целевой аудитории и конкретном жестовом языке", "Освещение, ракурс, скорость движения, частичное перекрытие руки", "Нагрузку на слабых устройствах при одновременном client/server inference", "UX-исследование с носителями жестового языка"], 7.18, 3.43, 4.75, 1.65, 13, bullet_color=ORANGE)
add_text(slide, "Вывод: прототип демонстрирует архитектурную траекторию, а не заменяет участие экспертов и валидацию словаря.", 0.82, 6.25, 11.55, 0.32, 14, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 13. Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
title(slide, "12. Roadmap", "Путь от демонстратора к полезному продукту", "Приоритеты идут от надёжности распознавания к расширению сценариев и масштабированию.", 13)
roadmap = [
    ("0–4 недели", "Базовая стабильность", ["Набор тестовых видео", "Latency/FPS telemetry", "UX-polish камеры"], CYAN),
    ("1–2 месяца", "Качество словаря", ["Датасет с носителями", "Калибровка confidence", "Языковая валидация"], VIOLET),
    ("2–4 месяца", "Productisation", ["Профили пользователей", "Сессии диалога", "Экспорт истории"], GREEN),
    ("4+ месяцев", "Scale", ["Mobile packaging", "3D avatar", "B2B/API integration"], ORANGE),
]
for i, (period, head, bullets, color) in enumerate(roadmap):
    x = 0.62 + i * 3.12
    card(slide, x, 2.58, 2.7, 3.25, color)
    add_text(slide, period, x + 0.25, 2.9, 2.1, 0.18, 10, color, bold=True)
    add_text(slide, head, x + 0.25, 3.34, 2.0, 0.42, 16, TEXT, bold=True)
    add_paragraphs(slide, bullets, x + 0.25, 4.18, 2.05, 1.1, 11, bullet_color=color)
    if i < 3:
        arrow(slide, x + 2.73, 4.22, x + 3.04, 4.22, color=LINE, width=2)
add_text(slide, "Фокус: сначала подтверждённая полезность и доверие, затем функциональная ширина и масштабирование.", 0.72, 6.35, 11.7, 0.3, 14, TEXT, bold=True, align=PP_ALIGN.CENTER)

# 14. Closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 14, "ФИНАЛ")
add_text(slide, "SignBridge", 0.68, 0.8, 6.0, 0.55, 30, TEXT, bold=True, font=FONT_DISPLAY)
add_text(slide, "Технология должна уменьшать барьеры,\nа не добавлять задержку между людьми.", 0.68, 1.72, 8.2, 0.82, 24, TEXT, bold=True, font=FONT_DISPLAY)
add_text(slide, "Следующий шаг: совместная валидация с носителями жестового языка и измерение качества в реальных сценариях.", 0.72, 2.9, 7.5, 0.46, 14, MUTED)
metric(slide, "Realtime", "контур руки в браузере", 0.92, 4.35, 2.0, CYAN)
metric(slide, "Local-first", "модель и обработка", 3.5, 4.35, 2.0, GREEN)
metric(slide, "Open API", "REST + WebSocket", 6.08, 4.35, 2.0, VIOLET)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.25, 1.15, 2.55, 3.7, SURFACE, LINE, True)
add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 10.12, 1.78, 0.78, 0.78, SURFACE_2, CYAN)
add_text(slide, "SB", 10.18, 2.03, 0.67, 0.18, 16, CYAN, bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER)
add_text(slide, "Спасибо", 9.68, 3.02, 1.7, 0.32, 22, TEXT, bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER)
add_text(slide, "Вопросы и обсуждение", 9.55, 3.57, 1.95, 0.25, 12, MUTED, align=PP_ALIGN.CENTER)
pill(slide, "SignBridge project", 9.63, 4.15, 1.6, CYAN)

prs.save(OUT)
print(OUT)
